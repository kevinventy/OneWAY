#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Devis One Way — feuille « Reçu de paiement » (facturation finale / quittance).

Ajoute une feuille reçu qui :
- reprend automatiquement les infos du devis (n°, client, total TTC) ;
- gère un règlement de 50 % (acompte ou solde) ou de 100 % (paiement intégral),
  ou un montant libre ;
- calcule le montant reçu, le reste à payer et le statut (soldé / en attente) ;
- écrit le MONTANT REÇU en toutes lettres (devise Ariary), automatiquement,
  via un 2ᵉ convertisseur ajouté à la feuille masquée « Lettres ».

Dépendances : pip install openpyxl Pillow
Usage :
    python3 scripts/devis-recu-paiement.py <source.xlsx> <sortie.xlsx>
"""
import io
import sys
import zipfile

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

DEVIS = "📋 Devis Client"
PARAMS = "⚙️ Paramètres"
HELPER = "Lettres"
RECU = "🧾 Reçu de paiement"

# Palette de la maison
DARK, NAVY, GRAY, INK, ORANGE = "001E3A5F", "002A4A72", "005B6B7F", "001F2937", "00E07B20"
LIGHT, BLUE, LINE = "00F4F7FB", "00EAF0F7", "00D9E1EC"
GREEN = "00128A4B"   # onglet du reçu (distinct du Calculateur orange)
MONEY = '#,##0" Ar"'

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def F(sz, color=INK, b=False, i=False):
    return Font(name="Calibri", size=sz, bold=b, italic=i, color=color)


def fill(c):
    return PatternFill("solid", fgColor=c)


def AL(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def put(ws, coord, value=None, font=None, fl=None, al=None, fmt=None, border=None):
    c = ws[coord]
    if value is not None:
        c.value = value
    if font:
        c.font = font
    if fl:
        c.fill = fl
    if al:
        c.alignment = al
    if fmt:
        c.number_format = fmt
    if border:
        c.border = border
    return c


def band(ws, rng, text, fl_color, sz=11):
    ws.merge_cells(rng)
    put(ws, rng.split(":")[0], text, F(sz, "00FFFFFF", b=True), fill(fl_color), AL("left"))


def label(ws, coord, text):
    put(ws, coord, text, F(9, GRAY, b=True), al=AL("left"))


def value(ws, rng, formula, money=False, big=False, onblue=False):
    ws.merge_cells(rng)
    tl = rng.split(":")[0]
    if onblue:
        put(ws, tl, formula, F(13 if big else 12, "00FFFFFF", b=True),
            fill(DARK), AL("right"), MONEY if money else None)
    else:
        put(ws, tl, formula, F(10, INK, b=big), al=AL("right"),
            fmt=MONEY if money else None)


def extend_helper(wb):
    """2e convertisseur (colonnes H/I) sur la feuille Lettres → montant reçu."""
    h = wb[HELPER]
    fin, mil = "$B$3:$B$1002", "$C$3:$C$1002"
    rows = {
        1: ("Reçu", f"='{RECU}'!$G$28"),
        2: ("n", "=INT(I1)"),
        3: ("g0", "=MOD(I2,1000)"),
        4: ("g1", "=MOD(INT(I2/1000),1000)"),
        5: ("g2", "=MOD(INT(I2/1000000),1000)"),
        6: ("g3", "=MOD(INT(I2/1000000000),1000)"),
        8: ("w0", f'=IF(I3=0,"",INDEX({fin},I3+1))'),
        9: ("w1", f'=IF(I4=0,"",IF(I4=1,"mille",INDEX({mil},I4+1)&" mille"))'),
        10: ("w2", f'=IF(I5=0,"",INDEX({fin},I5+1)&IF(I5>1," millions"," million"))'),
        11: ("w3", f'=IF(I6=0,"",INDEX({fin},I6+1)&IF(I6>1," milliards"," milliard"))'),
        13: ("lettres", '=IF(I2=0,"zéro",TRIM(I11&" "&I10&" "&I9&" "&I8))'),
    }
    for r, (lab, formula) in rows.items():
        h.cell(row=r, column=8, value=lab)
        h.cell(row=r, column=9, value=formula)


def build_recu(wb, src):
    if RECU in wb.sheetnames:
        del wb[RECU]
    ws = wb.create_sheet(RECU, index=4)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN

    widths = {"A": 3, "B": 26, "C": 16, "D": 12, "E": 12, "F": 18, "G": 18, "H": 14, "I": 14, "J": 3}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    dq = f"'{DEVIS}'"
    pq = f"'{PARAMS}'"
    lq = f"'{HELPER}'"

    # ── En-tête ────────────────────────────────────────────────────────────────
    for r in (2, 3, 4, 5, 6):
        ws.row_dimensions[r].height = 18
    ws.merge_cells("E3:I4")
    put(ws, "E3", "REÇU DE PAIEMENT", F(22, DARK, b=True), al=AL("center"))
    ws.merge_cells("E5:I6")
    put(ws, "E5", "Transport · Livraison · Suivi Digital", F(11, ORANGE), al=AL("center"))

    ws.row_dimensions[7].height = 6
    # Méta
    ws.row_dimensions[8].height = 20
    label(ws, "B8", "N° REÇU :")
    ws.merge_cells("C8:D8")
    put(ws, "C8", f'="REC-"&TEXT({pq}!C9,"0000")&"-"&TEXT(TODAY(),"MMYYYY")',
        F(10, DARK, b=True), al=AL("left"))
    label(ws, "E8", "DATE :")
    put(ws, "F8", "=TODAY()", F(10, INK, b=True), al=AL("left"), fmt="dd/mm/yyyy")
    label(ws, "G8", "RÉF. DEVIS :")
    ws.merge_cells("H8:I8")
    put(ws, "H8", f"={dq}!C12", F(10, INK, b=True), al=AL("right"))

    ws.row_dimensions[9].height = 6
    # Blocs client / émetteur
    ws.row_dimensions[10].height = 20
    band(ws, "B10:E10", "  👤  CLIENT", NAVY, sz=10)
    band(ws, "F10:I10", "  🚛  ONE WAY — TRANSPORT", ORANGE, sz=10)
    rows_info = [
        ("Nom / Société :", f"={dq}!C15", "Société :", f"={dq}!G15"),
        ("Adresse :",        f"={dq}!C16", "Adresse :", f"={dq}!G16"),
        ("Ville :",          f"={dq}!C17", "Téléphone :", f"={dq}!G17"),
        ("Téléphone :",      f"={dq}!C18", "Email :", f"={dq}!G18"),
        ("N° Client :",      f"={dq}!C20", "MVola :", f"={dq}!G19"),
    ]
    r = 11
    for lab_l, val_l, lab_r, val_r in rows_info:
        ws.row_dimensions[r].height = 18
        label(ws, f"B{r}", lab_l)
        ws.merge_cells(f"C{r}:E{r}")
        put(ws, f"C{r}", val_l, F(10, INK), al=AL("left"))
        label(ws, f"F{r}", lab_r)
        ws.merge_cells(f"G{r}:I{r}")
        put(ws, f"G{r}", val_r, F(10, INK), al=AL("left"))
        r += 1

    ws.row_dimensions[16].height = 6
    # Récapitulatif règlement
    ws.row_dimensions[17].height = 22
    band(ws, "B17:I17", "  💰  RÉCAPITULATIF DU RÈGLEMENT", DARK)
    recap = [
        (18, "Montant total du devis (TTC)", f"={dq}!G52"),
        (19, "Acompte — 50 % à la commande", f"=ROUND({dq}!G52*0.5,0)"),
        (20, "Solde — 50 % à la livraison", "=$G$18-$G$19"),
    ]
    for rr, lab_t, formula in recap:
        ws.row_dimensions[rr].height = 19
        ws.merge_cells(f"B{rr}:F{rr}")
        put(ws, f"B{rr}", lab_t, F(10, INK, b=(rr == 18)), fill(LIGHT), AL("right"))
        value(ws, f"G{rr}:I{rr}", formula, money=True)
        put(ws, f"G{rr}", None, fl=fill(LIGHT))

    ws.row_dimensions[21].height = 6
    # Paiement reçu
    ws.row_dimensions[22].height = 22
    band(ws, "B22:I22", "  🧾  PAIEMENT REÇU", DARK)

    def pay_row(rr, lab_t, formula=None, money=False, manual=None, fmt=None):
        ws.row_dimensions[rr].height = 19
        ws.merge_cells(f"B{rr}:F{rr}")
        put(ws, f"B{rr}", lab_t, F(10, GRAY, b=True), al=AL("right"))
        ws.merge_cells(f"G{rr}:I{rr}")
        if manual is not None:
            put(ws, f"G{rr}", manual, F(10, INK, b=True), fill(BLUE), AL("right"),
                fmt or (MONEY if money else None), BOX)
        else:
            put(ws, f"G{rr}", formula, F(10, INK), al=AL("right"),
                fmt=fmt or (MONEY if money else None))

    pay_row(23, "Nature du paiement :", manual="Acompte 50 % (à la commande)", fmt=None)
    pay_row(24, "Mode de paiement :", manual="MVola", fmt=None)
    pay_row(25, "Référence transaction :", manual="—", fmt=None)
    pay_row(26, "Montant libre (si nature = libre) :", manual=0, money=True)
    pay_row(27, "Déjà encaissé (paiements précédents) :", manual=0, money=True)

    # Montant reçu (mise en avant)
    ws.row_dimensions[28].height = 26
    ws.merge_cells("B28:F28")
    put(ws, "B28", "💵  MONTANT REÇU (Ar) :", F(12, "00FFFFFF", b=True), fill(DARK), AL("right"))
    value(ws, "G28:I28",
          f'=IF($G$23="Paiement intégral (100 %)",{dq}!G52,'
          f'IF($G$23="Solde 50 % (à la livraison)",{dq}!G52-ROUND({dq}!G52*0.5,0),'
          f'IF($G$23="Montant libre",$G$26,ROUND({dq}!G52*0.5,0))))',
          money=True, big=True, onblue=True)

    # Montant reçu en toutes lettres
    ws.row_dimensions[29].height = 24
    ws.merge_cells("B29:I29")
    put(ws, "B29", f'="Reçu la somme de "&{lq}!$I$13&" Ariary."',
        F(9.5, DARK, b=True, i=True), fill(LIGHT), AL("left", wrap=True), border=BOX)

    # Reste à payer + statut
    ws.row_dimensions[30].height = 19
    ws.merge_cells("B30:F30")
    put(ws, "B30", "Reste à payer après ce reçu :", F(10, INK, b=True), fill(LIGHT), AL("right"))
    value(ws, "G30:I30",
          f'=ROUND(IF($G$23="Acompte 50 % (à la commande)",{dq}!G52-$G$28,'
          f'IF($G$23="Montant libre",MAX({dq}!G52-$G$27-$G$28,0),0)),0)', money=True)
    put(ws, "G30", None, fl=fill(LIGHT))

    ws.row_dimensions[31].height = 20
    ws.merge_cells("B31:F31")
    put(ws, "B31", "Statut :", F(10, INK, b=True), al=AL("right"))
    ws.merge_cells("G31:I31")
    put(ws, "G31",
        '=IF($G$30<=0,"✅ SOLDÉ — Merci !","⏳ Reste à payer : "&TEXT($G$30,"#,##0")&" Ar")',
        F(11, DARK, b=True), al=AL("right"))

    ws.row_dimensions[32].height = 8
    # Mentions
    ws.row_dimensions[33].height = 28
    ws.merge_cells("B33:I33")
    put(ws, "B33",
        "Reçu valant quittance pour le montant ci-dessus, sous réserve d'encaissement. "
        "À conserver comme justificatif de paiement.",
        F(8.5, GRAY, i=True), al=AL("left", wrap=True))

    ws.row_dimensions[34].height = 8
    # Signatures
    ws.row_dimensions[35].height = 46
    ws.merge_cells("B35:E35")
    put(ws, "B35", "Le client", F(9, GRAY, b=True), al=AL("left", v="top"),
        border=Border(top=thin))
    ws.merge_cells("F35:I35")
    put(ws, "F35", "Pour One Way (cachet & signature)", F(9, GRAY, b=True),
        al=AL("left", v="top"), border=Border(top=thin))

    ws.row_dimensions[37].height = 22
    ws.merge_cells("B37:I37")
    put(ws, "B37",
        "One Way SARL  ·  Transport · Livraison · Suivi Digital  ·  Antananarivo, Madagascar 🇲🇬",
        F(9, GRAY, i=True), fill(LIGHT), AL("center"))

    # ── Listes déroulantes ──────────────────────────────────────────────────────
    dv_nat = DataValidation(
        type="list", allow_blank=True,
        formula1='"Acompte 50 % (à la commande),Solde 50 % (à la livraison),'
                 'Paiement intégral (100 %),Montant libre"')
    ws.add_data_validation(dv_nat)
    dv_nat.add(ws["G23"])
    dv_mode = DataValidation(type="list", allow_blank=True,
                             formula1='"MVola,Orange Money,Airtel Money,Virement,Espèces"')
    ws.add_data_validation(dv_mode)
    dv_mode.add(ws["G24"])

    # ── Logo ────────────────────────────────────────────────────────────────────
    with zipfile.ZipFile(src) as z:
        media = [m for m in z.namelist() if m.startswith("xl/media/")]
        if media:
            logo = XLImage(io.BytesIO(z.read(media[0])))
            logo.width = logo.height = 130
            ws.add_image(logo, "B2")

    ws.print_area = "A1:J37"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def main(src, dst):
    try:
        import PIL  # noqa: F401
    except ImportError:
        print("⚠️  Pillow absent : le logo risque d'être perdu (pip install Pillow).",
              file=sys.stderr)

    wb = openpyxl.load_workbook(src)
    if HELPER not in wb.sheetnames:
        sys.exit("La feuille « Lettres » est absente : lancez d'abord devis-montant-lettres.py.")

    build_recu(wb, src)
    extend_helper(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(dst)
    print(f"OK — feuille « {RECU} » créée (acompte 50 % / solde 50 % / 100 % / libre).")
    print("Montant reçu en toutes lettres branché sur la feuille masquée « Lettres ».")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: python3 scripts/devis-recu-paiement.py <source.xlsx> <sortie.xlsx>")
    main(sys.argv[1], sys.argv[2])
