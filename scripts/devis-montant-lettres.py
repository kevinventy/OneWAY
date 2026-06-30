#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Devis One Way — montant total en toutes lettres (automatique).

Ajoute, juste sous « TOTAL À PAYER (TTC) », une ligne qui écrit le montant total
EN TOUTES LETTRES (français, devise Ariary) et qui se met à jour automatiquement
quand le total change.

Comme Excel n'a pas de fonction native « nombre en lettres », la conversion est
faite uniquement par FORMULES, via une feuille auxiliaire masquée « Lettres »
contenant une table de conversion 0→999 (deux variantes pour gérer les pluriels
français « cent/cents » et « quatre-vingt/quatre-vingts »). Le devis ne contient
que des formules → tout reste dynamique, sans macro.

Dépendances : pip install openpyxl Pillow
Usage :
    python3 scripts/devis-montant-lettres.py <source.xlsx> <sortie.xlsx>
"""
import copy
import io
import sys
import zipfile

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import range_boundaries

SHEET = "📋 Devis Client"
HELPER = "Lettres"
TTC_CELL = "G52"          # cellule du TOTAL À PAYER (TTC)
INSERT_AT = 53            # nouvelle ligne (juste sous le TTC)
LAST_COL = 10             # A..J

# ── Générateur français 0→999 (deux formes : finale / avant « mille ») ──────────
UNITS = ["zéro", "un", "deux", "trois", "quatre", "cinq", "six", "sept", "huit",
         "neuf", "dix", "onze", "douze", "treize", "quatorze", "quinze", "seize",
         "dix-sept", "dix-huit", "dix-neuf"]
TENS = {20: "vingt", 30: "trente", 40: "quarante", 50: "cinquante", 60: "soixante"}


def two(n):                       # 0..99, forme finale
    if n < 20:
        return UNITS[n]
    if n < 70:
        t, u = (n // 10) * 10, n % 10
        if u == 0:
            return TENS[t]
        if u == 1 and t in (20, 30, 40, 50, 60):
            return TENS[t] + " et un"
        return TENS[t] + "-" + UNITS[u]
    if n < 80:                    # 70..79
        return "soixante et onze" if n == 71 else "soixante-" + UNITS[n - 60]
    if n == 80:
        return "quatre-vingts"
    return "quatre-vingt-" + UNITS[n - 80]   # 81..99


def two_nf(n):                    # 80 sans « s » (forme non finale)
    return "quatre-vingt" if n == 80 else two(n)


def three(n, mille=False):        # 0..999
    c, last2, parts = n // 100, n % 100, []
    if c == 1:
        parts.append("cent")
    elif c >= 2:
        parts.append(UNITS[c] + (" cents" if last2 == 0 and not mille else " cent"))
    if last2 > 0:
        parts.append(two_nf(last2) if mille else two(last2))
    return " ".join(parts)


def build_helper(wb):
    if HELPER in wb.sheetnames:
        del wb[HELPER]
    h = wb.create_sheet(HELPER)
    h.sheet_state = "hidden"
    h["A1"] = "Table de conversion (générée) — NE PAS MODIFIER"
    h["A2"], h["B2"], h["C2"] = "n", "final", "avant_mille"
    for n in range(1000):                       # A3:C1002  → n = 0..999
        r = n + 3
        h.cell(row=r, column=1, value=n)
        h.cell(row=r, column=2, value=three(n, mille=False))
        h.cell(row=r, column=3, value=three(n, mille=True))
    fin, mil = "$B$3:$B$1002", "$C$3:$C$1002"   # plages des deux formes

    # Cellules de contrôle (col. E libellé, F formule)
    ctl = {
        "E1": "Total", "F1": f"='{SHEET}'!$G$52",
        "E2": "n",  "F2": "=INT(F1)",
        "E3": "g0", "F3": "=MOD(F2,1000)",
        "E4": "g1", "F4": "=MOD(INT(F2/1000),1000)",
        "E5": "g2", "F5": "=MOD(INT(F2/1000000),1000)",
        "E6": "g3", "F6": "=MOD(INT(F2/1000000000),1000)",
        "E8":  "w0", "F8":  f"=IF(F3=0,\"\",INDEX({fin},F3+1))",
        "E9":  "w1", "F9":  f"=IF(F4=0,\"\",IF(F4=1,\"mille\",INDEX({mil},F4+1)&\" mille\"))",
        "E10": "w2", "F10": f"=IF(F5=0,\"\",INDEX({fin},F5+1)&IF(F5>1,\" millions\",\" million\"))",
        "E11": "w3", "F11": f"=IF(F6=0,\"\",INDEX({fin},F6+1)&IF(F6>1,\" milliards\",\" milliard\"))",
        "E13": "lettres", "F13": '=IF(F2=0,"zéro",TRIM(F11&" "&F10&" "&F9&" "&F8))',
    }
    for coord, val in ctl.items():
        h[coord] = val
    return f"'{HELPER}'!$F$13"


def insert_letters_row(ws, src, helper_ref):
    max_row = ws.max_row
    merges = [str(m) for m in ws.merged_cells.ranges]
    heights = {r: ws.row_dimensions[r].height for r in range(INSERT_AT, max_row + 1)
               if r in ws.row_dimensions and ws.row_dimensions[r].height is not None}

    for m in merges:
        ws.unmerge_cells(m)

    # décaler d'un cran vers le bas (de bas en haut)
    for r in range(max_row, INSERT_AT - 1, -1):
        for col in range(1, LAST_COL + 1):
            s = ws.cell(row=r, column=col)
            t = ws.cell(row=r + 1, column=col)
            t.value = s.value
            if s.has_style:
                t._style = copy.copy(s._style)
            s.value = None
            s._style = copy.copy(ws.cell(row=1, column=1)._style)

    # hauteurs
    for r in range(INSERT_AT, max_row + 2):
        if r in ws.row_dimensions:
            ws.row_dimensions[r].height = None
    for r, hh in heights.items():
        ws.row_dimensions[r + 1].height = hh
    ws.row_dimensions[INSERT_AT].height = 26

    # fusions (décaler celles >= point d'insertion)
    for m in merges:
        c1, r1, c2, r2 = range_boundaries(m)
        if r1 >= INSERT_AT:
            r1 += 1
            r2 += 1
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    ws.merge_cells(start_row=INSERT_AT, start_column=2, end_row=INSERT_AT, end_column=9)

    # contenu + style de la nouvelle ligne
    cell = ws.cell(row=INSERT_AT, column=2)
    cell.value = (f'="Arrêté le présent devis à la somme de "&{helper_ref}&" Ariary."')
    cell.font = Font(name="Calibri", size=9.5, italic=True, bold=True, color="001E3A5F")
    cell.fill = PatternFill("solid", fgColor="00F4F7FB")
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="00D9E1EC")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def main(src, dst):
    try:
        import PIL  # noqa: F401
    except ImportError:
        print("⚠️  Pillow absent : le logo risque d'être perdu (pip install Pillow).",
              file=sys.stderr)

    wb = openpyxl.load_workbook(src)
    ws = wb[SHEET]
    old_max = ws.max_row

    helper_ref = build_helper(wb)
    insert_letters_row(ws, src, helper_ref)

    ws.print_area = f"A2:J{old_max + 1}"

    # forcer le recalcul à l'ouverture (toutes les valeurs sont des formules)
    wb.calculation.fullCalcOnLoad = True

    if not ws._images:
        with zipfile.ZipFile(src) as z:
            media = [m for m in z.namelist() if m.startswith("xl/media/")]
            if media:
                logo = XLImage(io.BytesIO(z.read(media[0])))
                logo.width = logo.height = 150
                ws.add_image(logo, "B2")

    wb.save(dst)
    print(f"OK — ligne « montant en lettres » insérée en {INSERT_AT} (sous le TTC {TTC_CELL}).")
    print(f"Feuille auxiliaire masquée « {HELPER} » créée. Impression A2:J{old_max + 1}.")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: python3 scripts/devis-montant-lettres.py <source.xlsx> <sortie.xlsx>")
    main(sys.argv[1], sys.argv[2])
