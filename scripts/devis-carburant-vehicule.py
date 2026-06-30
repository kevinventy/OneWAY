#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Devis One Way — carburant automatique selon le véhicule + nettoyage.

1. Ajoute une LISTE DÉROULANTE de type de véhicule sur le champ « Type de
   véhicule » (C25), alimentée par la grille tarifaire (plage nommée Vehicules).
2. Recalcule automatiquement la ligne « Carburant » d'après le véhicule choisi :
   - quantité = litres de l'ALLER-RETOUR (distance × 2 × conso / 100) ;
   - prix au litre (gasoil ou essence) déduit de la grille → change tout seul.
   Le tarif/km de la ligne transport suit aussi le véhicule sélectionné.
3. Supprime la ligne « Retour véhicule à vide (repositionnement) » et recale
   formules de totaux, numérotation, fusions, hauteurs et zone d'impression.

Dépendances : pip install openpyxl Pillow
    (Pillow est indispensable, sinon openpyxl supprime le logo à l'enregistrement.)

Usage :
    python3 scripts/devis-carburant-vehicule.py <source.xlsx> <sortie.xlsx>
"""
import copy
import io
import sys
import zipfile

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

SHEET = "📋 Devis Client"
PARAMS = "⚙️ Paramètres"

VEH = f"'{PARAMS}'!$B$23:$B$32"          # noms des véhicules
CONSO = f"'{PARAMS}'!$E$23:$E$32"        # conso L/100km
CARB_KM = f"'{PARAMS}'!$F$23:$F$32"      # carburant Ar/km (gasoil/essence inclus)
TARIF_KM = f"'{PARAMS}'!$D$23:$D$32"     # tarif Ar/km hors carburant
MATCH = f"MATCH($C$25,{VEH},0)"

TRANSPORT_ROW = 31      # Transport marchandises — aller (tarif km)
FUEL_ROW = 32           # Carburant aller-retour
DELETE_ROW = 33         # Retour véhicule à vide (à supprimer)
FIRST_ITEM_ROW = 31
DEFAULT_VEHICLE = "🚚 Camion 5 tonnes"   # doit exister dans la grille (B26)
LAST_COL = 10           # A..J


def main(src, dst):
    try:
        import PIL  # noqa: F401
    except ImportError:
        print("⚠️  Pillow absent : le logo risque d'être perdu (pip install Pillow).",
              file=sys.stderr)

    wb = openpyxl.load_workbook(src)
    ws = wb[SHEET]
    max_row = ws.max_row

    # ── 1) Plage nommée + liste déroulante véhicule sur C25 ─────────────────────
    if "Vehicules" not in wb.defined_names:
        wb.defined_names.add(DefinedName("Vehicules", attr_text=VEH))
    ws["C25"] = DEFAULT_VEHICLE   # valeur qui correspond exactement à la grille

    # ── 2) Ligne transport : tarif/km selon véhicule, quantité = distance aller ──
    ws[f"F{TRANSPORT_ROW}"] = "=$C$24"
    ws[f"H{TRANSPORT_ROW}"] = f"=INDEX({TARIF_KM},{MATCH})"

    # ── 3) Ligne carburant : litres ALLER-RETOUR + prix/litre selon véhicule ────
    ws[f"C{FUEL_ROW}"] = "Carburant aller-retour (selon véhicule)"
    ws[f"E{FUEL_ROW}"] = "Carburant"
    ws[f"F{FUEL_ROW}"] = f"=ROUND($C$24*2*INDEX({CONSO},{MATCH})/100,0)"
    ws[f"G{FUEL_ROW}"] = "litre"
    # prix au litre (gasoil ou essence) = (carburant Ar/km) ÷ (conso L/100km) × 100
    ws[f"H{FUEL_ROW}"] = f"=ROUND(INDEX({CARB_KM},{MATCH})*100/INDEX({CONSO},{MATCH}),0)"
    ws[f"I{FUEL_ROW}"] = f"=F{FUEL_ROW}*H{FUEL_ROW}"

    # ── 4) Supprimer la ligne « Retour véhicule à vide » (décalage vers le haut) ─
    merges = [str(m) for m in ws.merged_cells.ranges]
    old_heights = {r: ws.row_dimensions[r].height
                   for r in range(DELETE_ROW, max_row + 1)
                   if r in ws.row_dimensions and ws.row_dimensions[r].height is not None}

    for m in merges:
        ws.unmerge_cells(m)

    # remonter les lignes 34..max d'un cran (de haut en bas : aucune source écrasée)
    default_style = copy.copy(ws.cell(row=1, column=1)._style)
    for r in range(DELETE_ROW + 1, max_row + 1):
        for col in range(1, LAST_COL + 1):
            s = ws.cell(row=r, column=col)
            t = ws.cell(row=r - 1, column=col)
            t.value = s.value
            if s.has_style:
                t._style = copy.copy(s._style)
    # vider l'ancienne dernière ligne
    for col in range(1, LAST_COL + 1):
        c = ws.cell(row=max_row, column=col)
        c.value = None
        c._style = copy.copy(default_style)

    # hauteurs de lignes
    for r in range(DELETE_ROW, max_row + 1):
        if r in ws.row_dimensions:
            ws.row_dimensions[r].height = None
    for r, h in old_heights.items():
        if r > DELETE_ROW:
            ws.row_dimensions[r - 1].height = h

    # fusions (la fusion de la ligne supprimée disparaît ; celles du dessous montent)
    for m in merges:
        min_col, min_row, max_col, max_r = range_boundaries(m)
        if min_row == DELETE_ROW:
            continue
        if min_row > DELETE_ROW:
            min_row -= 1
            max_r -= 1
        ws.merge_cells(start_row=min_row, start_column=min_col,
                       end_row=max_r, end_column=max_col)

    new_max = max_row - 1

    # ── 5) Recaler formules totaux + numérotation (lignes décalées de -1) ───────
    last_item = 47          # dernière prestation (17 lignes : 31..47)
    sub = last_item + 1     # 48 SOUS-TOTAL
    rem = sub + 1           # 49 Remise
    base = sub + 2          # 50 Base imposable
    tva = sub + 3           # 51 TVA
    ttc = sub + 4           # 52 TOTAL TTC
    ws.cell(row=sub,  column=7).value = f"=SUM(I{FIRST_ITEM_ROW}:I{last_item})"
    ws.cell(row=rem,  column=7).value = f"=-ROUND(G{sub}*F{rem}/100,0)"
    ws.cell(row=base, column=7).value = f"=G{sub}+G{rem}"
    ws.cell(row=tva,  column=7).value = f'=IF(F{tva}="Oui",ROUND(G{base}*0.2,0),0)'
    ws.cell(row=ttc,  column=7).value = f"=G{base}+G{tva}"

    # numérotation N° auto-incrémentée (1, puis =Bprec+1)
    ws.cell(row=FIRST_ITEM_ROW, column=2).value = 1
    for r in range(FIRST_ITEM_ROW + 1, last_item + 1):
        ws.cell(row=r, column=2).value = f"=B{r - 1}+1"

    # ── 6) Validations : véhicule (C25) + TVA Oui/Non (cellule TVA recalée) ─────
    ws.data_validations.dataValidation.clear()
    dv_veh = DataValidation(type="list", formula1="Vehicules", allow_blank=True,
                            error="Choisissez un type de véhicule dans la liste.",
                            errorTitle="Type de véhicule")
    ws.add_data_validation(dv_veh)
    dv_veh.add(ws["C25"])
    dv_tva = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    ws.add_data_validation(dv_tva)
    dv_tva.add(ws.cell(row=tva, column=6))   # F{tva}

    # ── 7) Zone d'impression (une ligne de moins) ──────────────────────────────
    ws.print_area = f"A2:J{new_max}"

    # ── 8) Sécurité logo ───────────────────────────────────────────────────────
    if not ws._images:
        with zipfile.ZipFile(src) as z:
            media = [m for m in z.namelist() if m.startswith("xl/media/")]
            if media:
                logo = XLImage(io.BytesIO(z.read(media[0])))
                logo.width = 150
                logo.height = 150
                ws.add_image(logo, "B2")

    wb.save(dst)
    print(f"OK — ligne « Retour véhicule à vide » supprimée. Prestations : "
          f"{FIRST_ITEM_ROW}..{last_item} (17 lignes).")
    print(f"Carburant ligne {FUEL_ROW} = aller-retour selon véhicule. "
          f"SOUS-TOTAL ligne {sub}, TTC ligne {ttc}. Impression A2:J{new_max}.")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: python3 scripts/devis-carburant-vehicule.py <source.xlsx> <sortie.xlsx>")
    main(sys.argv[1], sys.argv[2])
