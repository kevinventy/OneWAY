#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Devis One Way — suivi des paiements dans la feuille « Historique Devis ».

Étend le tableau de l'historique avec le suivi des règlements (acompte / solde) :
- colonnes ACOMPTE (Ar) et SOLDE (Ar) à saisir ;
- ENCAISSÉ, RESTE et PAIEMENT (statut) calculés AUTOMATIQUEMENT ;
- statistiques d'encaissement (total facturé / encaissé / reste / taux).

Le statut de paiement n'est calculé que pour les devis acceptés
(✅ Soldé / ⏳ Partiel / ❌ Impayé ; « — » sinon).

Dépendances : pip install openpyxl Pillow
Usage :
    python3 scripts/historique-paiements.py <source.xlsx> <sortie.xlsx>
"""
import copy
import sys

import openpyxl

SHEET = "📊 Historique Devis"
MONEY = '#,##0" Ar"'
FIRST, LAST = 7, 13          # lignes de données
ACC = '"✅ Accepté"'         # critère devis accepté

# Acompte / solde d'exemple (cohérents avec le statut du devis)
SAMPLE = {                    # ligne : (acompte, solde)
    7:  (1441500, 1441500),   # OW-0001 — soldé
    8:  (465500, 465500),     # OW-0002 — soldé
    9:  (1889500, None),      # OW-0003 — partiel
    10: (3289500, None),      # OW-0004 — partiel
    11: (None, None),         # OW-0005 — en attente
    12: (None, None),         # OW-0006 — refusé
    13: (None, None),         # OW-0007 — en attente
}


def clone(ws, src_coord, dst_coord):
    s, d = ws[src_coord], ws[dst_coord]
    if s.has_style:
        d._style = copy.copy(s._style)
    return d


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    ws = wb[SHEET]

    # ── Largeurs de colonnes + nouvelle marge ───────────────────────────────────
    for col in ("M", "N", "O", "P"):
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["Q"].width = 13
    ws.column_dimensions["R"].width = 3

    # ── En-têtes (ligne 6) ──────────────────────────────────────────────────────
    headers = {"M6": "ACOMPTE (Ar)", "N6": "SOLDE (Ar)", "O6": "ENCAISSÉ (Ar)",
               "P6": "RESTE (Ar)", "Q6": "PAIEMENT"}
    for coord, text in headers.items():
        clone(ws, "J6", coord).value = text

    # ── Lignes de données ───────────────────────────────────────────────────────
    for r in range(FIRST, LAST + 1):
        for col in ("M", "N", "O", "P"):       # style argent + zébrure de la ligne
            clone(ws, f"I{r}", f"{col}{r}")
        clone(ws, f"K{r}", f"Q{r}")            # style texte centré (statut)
        ac, so = SAMPLE.get(r, (None, None))
        ws[f"M{r}"].value = ac
        ws[f"N{r}"].value = so
        ws[f"O{r}"].value = f'=IF(K{r}<>"✅ Accepté","",N{r}+M{r})'
        ws[f"P{r}"].value = f'=IF(K{r}<>"✅ Accepté","",J{r}-M{r}-N{r})'
        ws[f"Q{r}"].value = (
            f'=IF(K{r}<>"✅ Accepté","—",'
            f'IF(M{r}+N{r}>=J{r},"✅ Soldé",'
            f'IF(M{r}+N{r}<=0,"❌ Impayé","⏳ Partiel")))')

    # ── Ligne TOTAL (14) : sommes sur les devis acceptés ────────────────────────
    for col in ("M", "N", "O", "P", "Q"):
        clone(ws, "I14", f"{col}14")
    ws["M14"].value = f"=SUMIF($K$7:$K$13,{ACC},M7:M13)"
    ws["N14"].value = f"=SUMIF($K$7:$K$13,{ACC},N7:N13)"
    ws["O14"].value = f"=SUMIF($K$7:$K$13,{ACC},O7:O13)"
    ws["P14"].value = f"=SUMIF($K$7:$K$13,{ACC},P7:P13)"

    # ── Bandeaux : étendre les fusions jusqu'à la colonne Q ─────────────────────
    for rng, new in (("B2:L4", "B2:Q4"), ("B16:L16", "B16:Q16")):
        if rng in [str(m) for m in ws.merged_cells.ranges]:
            ws.unmerge_cells(rng)
        ws.merge_cells(new)

    # ── Statistiques d'encaissement (lignes 22-26) ──────────────────────────────
    stats = [
        (22, "Total facturé (acceptés)", f"=SUMIF($K$7:$K$13,{ACC},J7:J13)", MONEY),
        (23, "Total encaissé",           "=SUM(O7:O13)", MONEY),
        (24, "Reste à encaisser",        f"=SUMIF($K$7:$K$13,{ACC},P7:P13)", MONEY),
        (25, "Taux d'encaissement",      "=IFERROR(C23/C22,0)", "0%"),
        (26, "Devis soldés / partiels",
         '=COUNTIF(Q7:Q13,"✅ Soldé")&" / "&COUNTIF(Q7:Q13,"⏳ Partiel")', "General"),
    ]
    for rr, lab, formula, fmt in stats:
        ws.row_dimensions[rr].height = 18.9
        clone(ws, "B17", f"B{rr}").value = lab
        c = clone(ws, "C17", f"C{rr}")
        c.value = formula
        c.number_format = fmt

    # ── Mise en page ────────────────────────────────────────────────────────────
    ws.print_area = "A1:R26"
    wb.calculation.fullCalcOnLoad = True
    wb.save(dst)
    print(f"OK — suivi des paiements ajouté à « {SHEET} » (colonnes M→Q).")
    print("ENCAISSÉ / RESTE / PAIEMENT automatiques + statistiques d'encaissement.")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: python3 scripts/historique-paiements.py <source.xlsx> <sortie.xlsx>")
    main(sys.argv[1], sys.argv[2])
